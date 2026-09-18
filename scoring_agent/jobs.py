"""Build the job list from an input path.

The user points the tool at a folder that contains:
  - one Excel file listing the module/lab names to search, in a column headed Module / Modules /
    Lab / Labs (or "Module Name" / "Lab Name") - capitalisation and spacing don't matter,
  - the scoring scripts, each a `.ps1` whose base name equals the module name, found anywhere
    under the folder (e.g. a "scripts"/"Scoring Scripts" subfolder).

A job is produced for every module name that has a matching `.ps1`.
"""
from __future__ import annotations
import os
from dataclasses import dataclass

from openpyxl import load_workbook

from .scriptmeta import parse_script, ScriptMeta

# Accepted column headings for the module/lab list. Matching is case-insensitive and ignores
# spaces, underscores and punctuation, so "Module", "MODULES", "Module Name", "Lab_Name",
# "Labs", "module-names" ... all work.
MODULE_HEADERS = {
    "module", "modules", "lab", "labs",
    "modulename", "modulenames", "modulesname", "modulesnames",
    "labname", "labnames", "labsname", "labsnames",
    "moduletitle", "labtitle", "moduleslist", "lablist", "labslist",
}
HEADER_SCAN_ROWS = 15  # the heading is not always on row 1 (title rows above it are common)


def _norm(v) -> str:
    """Normalise a header cell: lowercase, letters+digits only.

    '#' becomes 'num' first, so a "Lab #" column reads as 'labnum' and is NOT mistaken for the
    module column - that one holds numbers like 1.1.1, not module names.
    """
    if v is None:
        return ""
    return "".join(ch for ch in str(v).lower().replace("#", "num") if ch.isalnum())


def _module_column(row) -> int | None:
    if not row:
        return None
    for i, c in enumerate(row):
        if _norm(c) in MODULE_HEADERS:
            return i
    return None


@dataclass
class Job:
    module_name: str
    script_path: str
    meta: ScriptMeta


@dataclass
class ScanResult:
    jobs: list
    excel_path: str | None
    modules_without_script: list
    module_count: int


# Workbooks the agent itself wrote. They carry a "Module Name" column, so without this a
# previous run's report could be picked up as the module list for the next run.
OUR_OUTPUT = ("scoring-agent-report",)


def _find_excels(path: str) -> list:
    """Every workbook under the folder (a course folder often has an audit workbook too)."""
    if os.path.isfile(path) and path.lower().endswith((".xlsx", ".xlsm")):
        return [path]
    out = []
    for root, _dirs, files in os.walk(path):
        for f in sorted(files):
            low = f.lower()
            if not low.endswith((".xlsx", ".xlsm")) or f.startswith("~$"):
                continue
            if any(low.startswith(x) for x in OUR_OUTPUT):
                continue                      # our own report, not somebody's module list
            out.append(os.path.join(root, f))
    return out


def _read_modules(excel_path: str) -> list:
    wb = load_workbook(excel_path, read_only=True, data_only=True)
    modules = []
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        # Look for the heading row in the first few rows - sheets often carry a title above it.
        col = None
        for scanned, r in enumerate(rows, 1):
            col = _module_column(r)
            if col is not None or scanned >= HEADER_SCAN_ROWS:
                break
        if col is None:
            # no recognised header on this sheet; skip (another sheet may have it)
            continue
        for r in rows:
            if r is None or col >= len(r):
                continue
            v = r[col]
            if v is not None and str(v).strip():
                modules.append(str(v).strip())
        if modules:
            break  # first sheet with a Module/Lab column wins
    wb.close()
    # de-dup, preserve order
    seen, out = set(), []
    for m in modules:
        k = m.lower()
        if k not in seen:
            seen.add(k)
            out.append(m)
    return out


def _index_ps1(path: str) -> dict:
    out = {}
    base = path if os.path.isdir(path) else os.path.dirname(path)
    for root, _dirs, files in os.walk(base):
        for f in files:
            if f.lower().endswith(".ps1"):
                out.setdefault(os.path.splitext(f)[0].lower(), os.path.join(root, f))
    return out


def scan(path: str) -> ScanResult:
    excels = _find_excels(path)
    if not excels:
        raise FileNotFoundError("No Excel (.xlsx) with module names found in the folder.")

    ps1 = _index_ps1(path)
    # With more than one workbook present, the right one is the one whose names actually match
    # the scripts - so an audit/summary workbook sitting in the same folder can't hijack the run.
    best, best_modules, best_hits = None, [], -1
    for x in excels:
        try:
            mods = _read_modules(x)
        except Exception:
            continue
        hits = sum(1 for m in mods if m.lower() in ps1)
        if hits > best_hits or (hits == best_hits and best is None):
            best, best_modules, best_hits = x, mods, hits
        if hits and hits == len(mods):
            break                     # a perfect match - look no further

    excel, modules = best, best_modules
    if not modules:
        raise ValueError(
            "No module list found in the Excel. It needs a column headed Module / Modules / "
            "Lab / Labs (or 'Module Name' / 'Lab Name') - any capitalisation is fine.")

    jobs, missing = [], []
    for m in modules:
        script = ps1.get(m.lower())
        if not script:
            missing.append(m)
            continue
        jobs.append(Job(module_name=m, script_path=script, meta=parse_script(script)))
    return ScanResult(jobs=jobs, excel_path=excel, modules_without_script=missing,
                      module_count=len(modules))
