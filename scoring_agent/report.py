"""Write the end-of-run summary report as a formatted Excel file.

Columns: Module Name | Status | Version | VM (the VM the script targets / was made live on)
plus a Result/Notes column so failures are visible.
"""
from __future__ import annotations
import datetime as dt
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = "Segoe UI"
NAVY = "0B2A4A"
TEAL = "1FA5A0"
HDR_FILL = PatternFill("solid", fgColor=NAVY)
HDR_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
CELL = Font(name=FONT, size=10)
TITLE = Font(name=FONT, size=15, bold=True, color=NAVY)
SUB = Font(name=FONT, size=9, italic=True, color="5A6B7B")
BOLD = Font(name=FONT, size=10, bold=True)
THIN = Side(style="thin", color="D6DEE6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BAND = PatternFill("solid", fgColor="F3F7FA")

STATUS_FILL = {
    "Live": PatternFill("solid", fgColor="D8F3E6"),
    "InProgress": PatternFill("solid", fgColor="FFF4D6"),
    "Saved": PatternFill("solid", fgColor="FFF4D6"),
    "Failed": PatternFill("solid", fgColor="FBE0E0"),
}
STATUS_FONT = {
    "Live": Font(name=FONT, size=10, bold=True, color="1B7A4B"),
    "InProgress": Font(name=FONT, size=10, bold=True, color="9A6B00"),
    "Saved": Font(name=FONT, size=10, bold=True, color="9A6B00"),
    "Failed": Font(name=FONT, size=10, bold=True, color="B02020"),
}


def write_report(out_path: str, results: list, mode: str) -> str:
    """results: list of dicts {module, status, version, vm, note}."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    ws["A1"] = "Scoring Agent - Run Summary"
    ws["A1"].font = TITLE
    ok = sum(1 for r in results if r["status"] in ("Live", "InProgress", "Saved"))
    fail = sum(1 for r in results if r["status"] == "Failed")
    live = sum(1 for r in results if r["status"] == "Live")
    ws["A2"] = (f"{dt.datetime.now():%Y-%m-%d %H:%M}   |   Mode: {mode}   |   "
                f"{len(results)} modules   |   {live} live   |   {ok} ok   |   {fail} failed")
    ws["A2"].font = SUB

    headers = ["#", "Module Name", "Status", "Version", "VM", "Notes"]
    widths = [5, 60, 13, 9, 16, 46]
    HDR = 4
    for c, h in enumerate(headers, 1):
        cell = ws.cell(HDR, c, h)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER

    for i, r in enumerate(results):
        row = HDR + 1 + i
        vals = [i + 1, r["module"], r["status"], r.get("version", ""),
                r.get("vm", ""), r.get("note", "")]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row, c, v)
            cell.font = CELL
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=c in (2, 6),
                                       horizontal="center" if c in (1, 3, 4, 5) else "left")
            if i % 2:
                cell.fill = BAND
        st = r["status"]
        sc = ws.cell(row, 3)
        if st in STATUS_FILL:
            sc.fill = STATUS_FILL[st]
            sc.font = STATUS_FONT[st]

    last = HDR + len(results)
    tot = last + 1
    ws.cell(tot, 2, "TOTAL").font = BOLD
    ws.cell(tot, 3, f"{live} Live / {ok} ok / {fail} failed").font = BOLD
    for c in range(1, len(headers) + 1):
        ws.cell(tot, c).border = BORDER

    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[HDR].height = 22
    ws.freeze_panes = f"A{HDR + 1}"
    if results:
        ws.auto_filter.ref = f"A{HDR}:{get_column_letter(len(headers))}{last}"

    os.makedirs(os.path.dirname(out_path), exist_ok=True) if os.path.dirname(out_path) else None
    wb.save(out_path)
    return out_path
